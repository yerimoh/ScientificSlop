r"""Collect the score tables into one workbook so a later analysis or figure can start from a single file.

results/review_scores.xlsx
  iclr2026        one row per ICLR 2026 paper, raw scores and the three normalisations
  iclr_years      one row per ICLR 2017-2026 and FARS paper, same shape
  coverage        how many papers each system had scored, one row per snapshot
  README          what every column means and which normalisation a system supports

Raw values stay untouched. The normalised columns come from scripts/normalize_scores.py, where _def is the
definitional [0, 1] (only for scores whose definition fixes both ends), _mm is a robust min-max inside the corpus,
and _pct is the percentile rank inside the corpus. All three are oriented so that 1 means most AI-like.
"""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = f"{HERE}/results/review_scores.xlsx"

README = [
    ("What", "Per-paper score table used in the review-score analysis. Code review/scripts/{export_scores,normalize_scores,make_workbook}.py"),
    ("", ""),
    ("Sheet iclr2026", "ICLR 2026, a single year. One year, one scale, with detailed review scores and Pangram verdicts attached"),
    ("Sheet iclr_years", "ICLR 2017-2026 and FARS. The group column is FARS / reject / accept / oral"),
    ("Sheet coverage", "Number of papers scored per system per snapshot. Check here if measurement is still in progress"),
    ("", ""),
    ("Column <sys>", "Raw score. Units differ across systems, so do not compare directly"),
    ("Column <sys>_def", "[0,1] by definition. Our items and aggregates are failed units/checked units, so as-is; Pangram is an AI fraction, so as-is;"),
    ("", "reviewers are 1-10 ratings, so (10-x)/9. Detectors have no upper bound, so blank"),
    ("Column <sys>_mm", "min-max [0,1] clipped at the 1-99 percentile within the corpus. Present for all systems, but it means position within this corpus"),
    ("Column <sys>_pct", "Percentile rank [0,1] within the corpus. The value the figures use"),
    ("Direction", "For all three normalizations, closer to 1 means more AI-like. Binoculars and reviewers are AI when the raw score is low, so their sign is flipped"),
    ("", ""),
    ("Column rating", "Mean public review score. Scales differ by year, so compare only within a year"),
    ("Column soundness/presentation/contribution/confidence", "Detailed review scores. 2026 only"),
    ("Column pangram_fraction_ai", "AI fraction published by Pangram. Not run by us; the public ICLR 2026 verdict"),
    ("Column <item>_weak", "Flag that the item's denominator is below the minimum. If 1, exclude from aggregation"),
    ("Column <item>_denominator", "Number of units the item actually checked"),
    ("Column quality_excluded", "Paper caught by the extraction-quality rules. If 1, excluded from analysis"),
]


def sheet_from_csv(wb, name, path, freeze="B2"):
    if not os.path.exists(path):
        return 0
    ws = wb.create_sheet(name)
    with open(path) as fh:
        for i, row in enumerate(csv.reader(fh)):
            ws.append(row if i == 0 else [try_num(v) for v in row])
    for c in range(1, ws.max_column + 1):
        ws.cell(row=1, column=c).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(c)].width = max(10, min(22, len(str(ws.cell(row=1, column=c).value or "")) + 3))
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    return ws.max_row - 1


def try_num(v):
    if v in ("", None):
        return None
    try:
        f = float(v)
        return int(f) if f.is_integer() and abs(f) < 1e15 and "." not in v else f
    except ValueError:
        return v


wb = Workbook()
wb.remove(wb.active)
ws = wb.create_sheet("README")
ws.append(["Item", "Description"])
for a, b in README:
    ws.append([a, b])
ws.column_dimensions["A"].width = 44
ws.column_dimensions["B"].width = 110
for r in range(1, ws.max_row + 1):
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
ws.cell(row=1, column=1).font = Font(bold=True); ws.cell(row=1, column=2).font = Font(bold=True)
n26 = sheet_from_csv(wb, "iclr2026", f"{HERE}/results/scores_2026_norm.csv")
ny = sheet_from_csv(wb, "iclr_years", f"{HERE}/results/scores_years_norm.csv")
nc = sheet_from_csv(wb, "coverage", f"{HERE}/results/coverage.csv", freeze="A2")
tmp = OUT + ".tmp"
wb.save(tmp)
os.replace(tmp, OUT)
print(f"review_scores.xlsx  iclr2026 {n26} rows, iclr_years {ny} rows, coverage {nc} rows")
